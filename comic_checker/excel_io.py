from openpyxl import load_workbook


def open_workbook(path):
    return load_workbook(path)


def find_column(ws, header_name, header_row=1):
    """ヘッダー名から列番号を探す。見つからなければ例外を送出する。"""
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=header_row, column=col).value == header_name:
            return col
    raise ValueError(f"列「{header_name}」が見つかりません（{header_row}行目のヘッダーを確認してください）")


def ensure_column(ws, header_name, header_row=1):
    """ヘッダー名から列番号を探す。見つからなければ末尾に新規作成して列番号を返す。"""
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=header_row, column=col).value == header_name:
            return col
    new_col = ws.max_column + 1
    ws.cell(row=header_row, column=new_col, value=header_name)
    return new_col
