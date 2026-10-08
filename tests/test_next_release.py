import pytest
from openpyxl import Workbook, load_workbook

from comic_checker import next_release, rakuten_books_client
from comic_checker.rakuten_books_client import ReleaseInfo

KEY_VALUE_ID = "dummy-app-id-value"
KEY_VALUE_ACCESS = "dummy-access-key-value"

COMPLETION_HEADERS = ["連載状況", "判定方法", "判定根拠", "参照タイトル", "参照URL"]
OUT_HEADERS = ["次巻発売日", "次巻タイトル", "最新既刊", "発売日参照URL"]


@pytest.fixture(autouse=True)
def _keys(monkeypatch):
    monkeypatch.setenv("RAKUTEN_APP_ID", KEY_VALUE_ID)
    monkeypatch.setenv("RAKUTEN_ACCESS_KEY", KEY_VALUE_ACCESS)


def _make_excel(tmp_path, titles, extra_headers=(), extra_values=None):
    path = tmp_path / "list.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["タイトル", *extra_headers])
    for i, title in enumerate(titles):
        row = [title]
        if extra_values:
            row += list(extra_values[i])
        ws.append(row)
    wb.save(path)
    return path


def _read_row(path, row):
    ws = load_workbook(path).active
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    return {h: ws.cell(row=row, column=c).value for h, c in headers.items()}


def _patch_find(monkeypatch, mapping):
    calls = []

    def fake(title, today=None):
        calls.append(title)
        value = mapping[title]
        if isinstance(value, Exception):
            raise value
        return value

    monkeypatch.setattr(rakuten_books_client, "find_next_release", fake)
    return calls


def _next(**kw):
    base = dict(
        query="q",
        status="次巻あり",
        next_title="架空の物語 3巻",
        next_sales_date="2030年1月5日",
        latest_title="架空の物語 2巻",
        latest_sales_date="2029年10月5日",
        url="https://example.invalid/3",
    )
    base.update(kw)
    return ReleaseInfo(**base)


def test_status_next(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["架空の物語"])
    _patch_find(monkeypatch, {"架空の物語": _next()})
    next_release.run(path)
    row = _read_row(path, 2)
    assert row["次巻発売日"] == "2030年1月5日"
    assert row["次巻タイトル"] == "架空の物語 3巻"
    assert row["最新既刊"] == "架空の物語 2巻（2029年10月5日）"
    assert row["発売日参照URL"] == "https://example.invalid/3"


def test_special_edition_note(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["架空の物語"])
    _patch_find(monkeypatch, {"架空の物語": _next(next_has_special_edition=True)})
    next_release.run(path)
    assert _read_row(path, 2)["次巻タイトル"] == "架空の物語 3巻（特装版あり）"


def test_latest_text_variants(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["甲", "乙"])
    _patch_find(
        monkeypatch,
        {
            "甲": _next(latest_sales_date=""),
            "乙": _next(latest_title="", latest_sales_date="2029年10月5日"),
        },
    )
    next_release.run(path)
    assert _read_row(path, 2)["最新既刊"] == "架空の物語 2巻"
    assert _read_row(path, 3)["最新既刊"] in (None, "")


def test_status_undecided(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["架空の物語"])
    info = _next(status="未定", next_title="", next_sales_date="")
    _patch_find(monkeypatch, {"架空の物語": info})
    next_release.run(path)
    row = _read_row(path, 2)
    assert row["次巻発売日"] == "未定"
    assert row["次巻タイトル"] in (None, "")
    assert row["最新既刊"] == "架空の物語 2巻（2029年10月5日）"
    assert row["発売日参照URL"] == "https://example.invalid/3"


@pytest.mark.parametrize("status", ["要確認", "ヒットなし", "エラー"])
def test_other_statuses(tmp_path, monkeypatch, status):
    path = _make_excel(tmp_path, ["架空の物語"])
    info = ReleaseInfo(
        query="q", status=status, latest_title="残り", latest_sales_date="d", url="u"
    )
    _patch_find(monkeypatch, {"架空の物語": info})
    next_release.run(path)
    row = _read_row(path, 2)
    assert row["次巻発売日"] == status
    for header in OUT_HEADERS[1:]:
        assert row[header] in (None, "")


def test_overwrite_clears_previous_values(tmp_path, monkeypatch):
    path = _make_excel(
        tmp_path,
        ["架空の物語"],
        extra_headers=OUT_HEADERS,
        extra_values=[["古い日付", "古い次巻", "古い既刊", "古いURL"]],
    )
    _patch_find(monkeypatch, {"架空の物語": ReleaseInfo(query="q", status="ヒットなし")})
    next_release.run(path)
    row = _read_row(path, 2)
    assert row["次巻発売日"] == "ヒットなし"
    for header in OUT_HEADERS[1:]:
        assert row[header] in (None, "")


def test_empty_title_rows_are_skipped(tmp_path, monkeypatch):
    path = _make_excel(
        tmp_path,
        [None, "   ", "架空の物語"],
        extra_headers=OUT_HEADERS,
        extra_values=[["a", "b", "c", "d"], ["e", "f", "g", "h"], ["", "", "", ""]],
    )
    calls = _patch_find(monkeypatch, {"架空の物語": _next()})
    next_release.run(path)
    assert calls == ["架空の物語"]
    assert [_read_row(path, 2)[h] for h in OUT_HEADERS] == ["a", "b", "c", "d"]
    assert [_read_row(path, 3)[h] for h in OUT_HEADERS] == ["e", "f", "g", "h"]
    assert _read_row(path, 4)["次巻発売日"] == "2030年1月5日"


def test_title_is_stripped_and_str(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["  架空の物語 ", 123])
    calls = _patch_find(
        monkeypatch,
        {"架空の物語": _next(), "123": ReleaseInfo(query="q", status="ヒットなし")},
    )
    next_release.run(path)
    assert calls == ["架空の物語", "123"]


def test_columns_created_at_end(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["架空の物語"], extra_headers=["備考"], extra_values=[["メモ"]])
    _patch_find(monkeypatch, {"架空の物語": _next()})
    next_release.run(path)
    ws = load_workbook(path).active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers == ["タイトル", "備考", *OUT_HEADERS]
    assert _read_row(path, 2)["備考"] == "メモ"


def test_missing_title_column_raises(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["架空の物語"])
    _patch_find(monkeypatch, {})
    with pytest.raises(ValueError):
        next_release.run(path, title_column="存在しない列")


def test_completion_columns_untouched(tmp_path, monkeypatch):
    values = [["連載中", "Wikipedia", "根拠", "参照題", "https://example.invalid/w"]]
    path = _make_excel(tmp_path, ["架空の物語"], extra_headers=COMPLETION_HEADERS, extra_values=values)
    _patch_find(monkeypatch, {"架空の物語": _next()})
    next_release.run(path)
    row = _read_row(path, 2)
    assert [row[h] for h in COMPLETION_HEADERS] == values[0]


def test_row_exception_becomes_error_and_continues(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["甲", "乙"])
    _patch_find(monkeypatch, {"甲": ValueError("boom"), "乙": _next()})
    events = []
    next_release.run(
        path, on_progress=lambda *a, error=None: events.append((a[0], a[3], error))
    )
    first = _read_row(path, 2)
    assert first["次巻発売日"] == "エラー"
    for header in OUT_HEADERS[1:]:
        assert first[header] in (None, "")
    assert _read_row(path, 3)["次巻発売日"] == "2030年1月5日"
    assert isinstance(events[0][2], ValueError)
    assert events[1][2] is None


def test_progress_called_per_row(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["甲", "乙"])
    _patch_find(monkeypatch, {"甲": _next(), "乙": ReleaseInfo(query="q", status="要確認")})
    events = []
    n = next_release.run(path, on_progress=lambda *a, error=None: events.append(a))
    assert n == 2
    assert events == [
        (1, 2, "甲", "次巻あり", "2030年1月5日"),
        (2, 2, "乙", "要確認", "要確認"),
    ]


def test_save_every(tmp_path, monkeypatch):
    path = _make_excel(tmp_path, ["甲", "乙", "丙"])
    _patch_find(monkeypatch, {t: _next() for t in ["甲", "乙", "丙"]})
    snapshots = []

    def progress(i, total, title, status, nxt, error=None):
        if i == 3:
            snapshots.append(_read_row(path, 2)["次巻発売日"])

    next_release.run(path, save_every=2, on_progress=progress)
    # 2件目の時点で保存されているため、3件目の処理中に1行目の結果が読める
    assert snapshots == ["2030年1月5日"]


def test_no_rows_message(tmp_path, monkeypatch, capsys):
    path = _make_excel(tmp_path, [None])
    _patch_find(monkeypatch, {})
    assert next_release.run(path) == 0
    assert "処理対象の行がありません（タイトル未入力）" in capsys.readouterr().out


def test_run_without_keys_raises_before_opening_excel(tmp_path, monkeypatch):
    monkeypatch.delenv("RAKUTEN_APP_ID")
    with pytest.raises(RuntimeError):
        next_release.run(tmp_path / "does_not_exist.xlsx")


def test_main_without_keys_returns_nonzero(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("RAKUTEN_ACCESS_KEY")
    path = _make_excel(tmp_path, ["架空の物語"])
    before = path.read_bytes()
    _patch_find(monkeypatch, {"架空の物語": _next()})
    assert next_release.main([str(path)]) == 1
    assert path.read_bytes() == before
    captured = capsys.readouterr()
    assert "RAKUTEN_ACCESS_KEY" in captured.err
    assert KEY_VALUE_ID not in captured.out + captured.err
    assert KEY_VALUE_ACCESS not in captured.out + captured.err


def test_main_missing_file(tmp_path, capsys):
    assert next_release.main([str(tmp_path / "none.xlsx")]) == 1
    assert "ファイルが見つかりません" in capsys.readouterr().err


def test_main_progress_output_and_error_hides_message(tmp_path, monkeypatch, capsys):
    path = _make_excel(tmp_path, ["甲", "乙"])
    secret = "https://example.invalid/?accessKey=SECRETVALUE"
    _patch_find(monkeypatch, {"甲": _next(), "乙": RuntimeError(secret)})
    assert next_release.main([str(path)]) == 0
    captured = capsys.readouterr()
    assert "2件を処理します" in captured.out
    assert "[1/2] 甲 -> 次巻あり (2030年1月5日)" in captured.out
    assert "完了しました" in captured.out
    assert "[2/2] 乙 -> エラー RuntimeError" in captured.err
    assert "SECRETVALUE" not in captured.out + captured.err
    assert "Traceback" not in captured.err
    assert _read_row(path, 3)["次巻発売日"] == "エラー"


def test_main_sheet_and_header_row(tmp_path, monkeypatch):
    path = tmp_path / "list.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "対象"
    ws.append(["説明行"])
    ws.append(["作品名欄"])
    ws.append(["架空の物語"])
    wb.save(path)
    _patch_find(monkeypatch, {"架空の物語": _next()})
    rc = next_release.main(
        [str(path), "--sheet", "対象", "--header-row", "2", "--title-column", "作品名欄"]
    )
    assert rc == 0
    ws = load_workbook(path)["対象"]
    assert ws.cell(row=2, column=2).value == "次巻発売日"
    assert ws.cell(row=3, column=2).value == "2030年1月5日"
