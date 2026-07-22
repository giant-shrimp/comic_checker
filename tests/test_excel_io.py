import pytest
from openpyxl import Workbook

from comic_checker.excel_io import ensure_column, find_column


def make_ws():
    wb = Workbook()
    ws = wb.active
    ws.append(["タイトル", "連載状況"])
    ws.append(["ONE PIECE", None])
    return ws


def test_find_column_returns_existing_index():
    ws = make_ws()
    assert find_column(ws, "タイトル") == 1
    assert find_column(ws, "連載状況") == 2


def test_find_column_raises_when_missing():
    ws = make_ws()
    with pytest.raises(ValueError):
        find_column(ws, "存在しない列")


def test_ensure_column_returns_existing_index_without_adding():
    ws = make_ws()
    before = ws.max_column
    assert ensure_column(ws, "連載状況") == 2
    assert ws.max_column == before


def test_ensure_column_creates_new_column_when_missing():
    ws = make_ws()
    before = ws.max_column
    col = ensure_column(ws, "マンバURL")
    assert col == before + 1
    assert ws.cell(row=1, column=col).value == "マンバURL"
