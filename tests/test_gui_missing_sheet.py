"""保存したシート名が Excel に無いときの案内のテスト（Issue #551）。名前はすべて架空。"""
import tkinter as tk
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from comic_checker import gui, user_config
from comic_checker.gui import App, MISSING_SHEET_MESSAGE, find_missing_sheet
from comic_checker.main import RowCounts

EXPECTED_MESSAGE = (
    "シート「存在しないシート」が見つかりません。"
    "「設定」の画面のシート名を確かめてください。"
)


@pytest.fixture
def app():
    root = tk.Tk()
    widget = App(master=root)
    widget.pack(fill="both", expand=True)
    root.update()
    yield widget
    widget._stop_elapsed()
    root.destroy()


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "list.xlsx"
    wb = Workbook()
    wb.active.title = "シートA"
    wb.active.append(["タイトル"])
    wb.create_sheet("シートB").append(["タイトル"])
    wb.save(path)
    return str(path)


def save_sheet(sheet):
    user_config.save_values({gui.KEY_SHEET: sheet})


def invoke(app, method_name):
    """指定の入口を呼び、(Thread のモック, count_rows のモック) を返す。"""
    with patch(
        "comic_checker.gui.count_rows", return_value=RowCounts(pending=3, filled=0)
    ) as count:
        with patch("comic_checker.gui.threading.Thread") as thread_cls:
            getattr(app, method_name)()
    return thread_cls, count


def log_messages(app):
    return [(level, message) for _time, level, message in app._log_entries]


ENTRY_POINTS = pytest.mark.parametrize("method_name", ["_on_run", "_on_next_release"])


def test_message_text_is_fixed():
    assert MISSING_SHEET_MESSAGE.format(sheet="存在しないシート") == EXPECTED_MESSAGE


@ENTRY_POINTS
def test_missing_sheet_stops_before_starting(app, book, method_name):
    save_sheet("存在しないシート")
    app.excel_path.set(book)
    thread_cls, count = invoke(app, method_name)
    thread_cls.assert_not_called()
    count.assert_not_called()
    assert ("ERROR", EXPECTED_MESSAGE) in log_messages(app)
    assert app.status_text.get() == EXPECTED_MESSAGE
    assert str(app.run_button.cget("state")) == "normal"
    assert str(app.next_button.cget("state")) == "normal"


@ENTRY_POINTS
def test_existing_sheet_starts_worker_with_sheet(app, book, method_name):
    save_sheet("シートB")
    app.excel_path.set(book)
    thread_cls, _count = invoke(app, method_name)
    thread_cls.assert_called_once()
    assert thread_cls.call_args.kwargs["kwargs"]["sheet"] == "シートB"
    assert EXPECTED_MESSAGE not in [m for _level, m in log_messages(app)]


@ENTRY_POINTS
def test_no_saved_sheet_does_not_open_workbook(app, book, method_name):
    app.excel_path.set(book)
    # シートの確認だけを見る（列の確認はブックを開くため止める。Issue #553）。
    with patch("comic_checker.gui.find_missing_title_column", return_value=False):
        with patch("comic_checker.gui.load_workbook") as loader:
            thread_cls, _count = invoke(app, method_name)
    loader.assert_not_called()
    thread_cls.assert_called_once()


@ENTRY_POINTS
def test_unopenable_file_falls_through_without_message(app, tmp_path, method_name):
    save_sheet("存在しないシート")
    app.excel_path.set(str(tmp_path / "no_such_file.xlsx"))
    thread_cls, _count = invoke(app, method_name)
    thread_cls.assert_called_once()
    assert all(
        "見つかりません" not in message for _level, message in log_messages(app)
    )


@pytest.mark.parametrize(
    "sheet, expected", [("存在しないシート", True), ("シートA", False)]
)
def test_find_missing_sheet_opens_read_only_and_closes_once(sheet, expected):
    wb = MagicMock()
    wb.sheetnames = ["シートA", "シートB"]
    with patch("comic_checker.gui.load_workbook", return_value=wb) as loader:
        assert find_missing_sheet("/tmp/list.xlsx", sheet) is expected
    loader.assert_called_once_with("/tmp/list.xlsx", read_only=True)
    wb.close.assert_called_once_with()


@pytest.mark.parametrize("sheet", [None, ""])
def test_find_missing_sheet_skips_when_no_sheet(sheet):
    with patch("comic_checker.gui.load_workbook") as loader:
        assert find_missing_sheet("/tmp/list.xlsx", sheet) is False
    loader.assert_not_called()


def test_find_missing_sheet_returns_false_when_open_fails():
    with patch("comic_checker.gui.load_workbook", side_effect=OSError("x")):
        assert find_missing_sheet("/tmp/list.xlsx", "シートA") is False
