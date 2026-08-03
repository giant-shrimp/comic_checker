from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from comic_checker.main import run
from comic_checker.manba_client import LookupResult
from comic_checker.wikipedia_client import WikipediaJudgment


def make_workbook(tmp_path, titles):
    wb = Workbook()
    ws = wb.active
    ws.append(["タイトル", "連載状況"])
    for title in titles:
        ws.append([title, None])
    path = tmp_path / "test.xlsx"
    wb.save(path)
    return path


def make_fake_manba_client(lookup_side_effect):
    fake_client = MagicMock()
    fake_client.lookup.side_effect = lookup_side_effect
    fake_client.__enter__.return_value = fake_client
    fake_client.__exit__.return_value = False
    return fake_client


def read_header(ws):
    return [cell.value for cell in ws[1]]


def read_row(ws, row):
    return {header: ws.cell(row=row, column=col).value for col, header in enumerate(read_header(ws), start=1)}


def test_run_uses_wikipedia_and_never_touches_manba_by_default(tmp_path):
    excel_path = make_workbook(tmp_path, ["ONE PIECE", "進撃の巨人"])

    wiki_results = [
        WikipediaJudgment(
            query="ONE PIECE",
            status="連載中",
            matched_title="ONE PIECE",
            url="https://ja.wikipedia.org/wiki/ONE_PIECE",
            evidence="1997年34号から連載中。",
        ),
        WikipediaJudgment(
            query="進撃の巨人",
            status="完結",
            matched_title="進撃の巨人",
            url="https://ja.wikipedia.org/wiki/進撃の巨人",
            evidence="5月号をもって完結した。",
        ),
    ]

    calls = []

    def on_progress(index, total, title, status, matched_title, error=None):
        calls.append((index, total, title, status, matched_title, error))

    with patch("comic_checker.main.wikipedia_client.judge", side_effect=wiki_results), patch(
        "comic_checker.manba_client.ManbaClient"
    ) as mock_manba_client_cls:
        total = run(excel_path, on_progress=on_progress)

    assert total == 2
    mock_manba_client_cls.assert_not_called()
    assert calls == [
        (1, 2, "ONE PIECE", "連載中", "ONE PIECE", None),
        (2, 2, "進撃の巨人", "完結", "進撃の巨人", None),
    ]

    wb = _reload(excel_path)
    ws = wb.active
    row1 = read_row(ws, 2)
    assert row1["連載状況"] == "連載中"
    assert row1["判定方法"] == "Wikipedia"
    assert row1["判定根拠"] == "1997年34号から連載中。"
    assert row1["参照タイトル"] == "ONE PIECE"
    assert row1["参照URL"] == "https://ja.wikipedia.org/wiki/ONE_PIECE"


def test_run_records_unknown_status_without_fallback_by_default(tmp_path):
    excel_path = make_workbook(tmp_path, ["謎の作品"])

    wiki_result = WikipediaJudgment(query="謎の作品", status="判定不能")

    with patch("comic_checker.main.wikipedia_client.judge", return_value=wiki_result), patch(
        "comic_checker.manba_client.ManbaClient"
    ) as mock_manba_client_cls:
        total = run(excel_path)

    assert total == 1
    mock_manba_client_cls.assert_not_called()

    wb = _reload(excel_path)
    ws = wb.active
    row1 = read_row(ws, 2)
    assert row1["連載状況"] == "判定不能"
    assert row1["判定方法"] == "Wikipedia"


def test_run_falls_back_to_manba_when_enabled_and_reuses_client(tmp_path):
    excel_path = make_workbook(tmp_path, ["謎の作品1", "謎の作品2"])

    wiki_results = [
        WikipediaJudgment(query="謎の作品1", status="判定不能"),
        WikipediaJudgment(query="謎の作品2", status="記事なし"),
    ]
    fake_manba_client = make_fake_manba_client(
        [
            LookupResult(
                query="謎の作品1",
                matched_title="謎の作品1（マンバ）",
                board_id="1",
                url="https://manba.co.jp/boards/1",
                status="連載中",
            ),
            LookupResult(
                query="謎の作品2",
                matched_title="謎の作品2（マンバ）",
                board_id="2",
                url="https://manba.co.jp/boards/2",
                status="完結",
            ),
        ]
    )

    with patch("comic_checker.main.wikipedia_client.judge", side_effect=wiki_results), patch(
        "comic_checker.manba_client.ManbaClient", return_value=fake_manba_client
    ) as mock_manba_client_cls:
        total = run(excel_path, fallback_to_manba=True)

    assert total == 2
    # 遅延生成: 1行目のフォールバック発生時に一度だけ生成され、2行目は使い回される。
    mock_manba_client_cls.assert_called_once()
    fake_manba_client.__exit__.assert_called_once()

    wb = _reload(excel_path)
    ws = wb.active
    row1 = read_row(ws, 2)
    row2 = read_row(ws, 3)
    assert row1["連載状況"] == "連載中"
    assert row1["判定方法"] == "manba.co.jp"
    assert row1["参照タイトル"] == "謎の作品1（マンバ）"
    assert row2["連載状況"] == "完結"
    assert row2["判定方法"] == "manba.co.jp"


def test_run_fails_fast_when_fallback_requested_but_manba_client_unimportable(tmp_path):
    # 軽量ビルド（playwright非同梱）を想定した回帰テスト: --fallback-to-manba指定時、
    # manba_clientのimportに失敗したら、行処理を始める前に明確なRuntimeErrorで
    # 即座に失敗する（1行ごとに「エラー」を積み重ねる形にはならない）。
    excel_path = make_workbook(tmp_path, ["謎の作品"])

    with patch(
        "comic_checker.main._import_manba_client",
        side_effect=RuntimeError("manba.co.jpフォールバック（--fallback-to-manba）にはplaywrightが必要です"),
    ):
        with pytest.raises(RuntimeError, match="playwright"):
            run(excel_path, fallback_to_manba=True)

    # 行処理そのものに到達していないこと（Wikipedia判定すら呼ばれない）を確認する。
    wb = _reload(excel_path)
    ws = wb.active
    assert ws.cell(row=2, column=2).value is None


def test_run_passes_error_to_on_progress_and_continues(tmp_path):
    excel_path = make_workbook(tmp_path, ["ONE PIECE", "NARUTO"])
    boom = RuntimeError("boom")

    with patch(
        "comic_checker.main.wikipedia_client.judge",
        side_effect=[boom, WikipediaJudgment(query="NARUTO", status="完結", matched_title="NARUTO")],
    ):
        calls = []

        def on_progress(index, total, title, status, matched_title, error=None):
            calls.append((index, total, title, status, matched_title, error))

        total = run(excel_path, on_progress=on_progress)

    assert total == 2
    assert calls[0] == (1, 2, "ONE PIECE", "エラー", "", boom)
    assert calls[1] == (2, 2, "NARUTO", "完結", "NARUTO", None)


def test_run_without_on_progress_does_not_raise(tmp_path):
    excel_path = make_workbook(tmp_path, ["ONE PIECE"])

    with patch(
        "comic_checker.main.wikipedia_client.judge",
        return_value=WikipediaJudgment(query="ONE PIECE", status="連載中", matched_title="ONE PIECE"),
    ):
        total = run(excel_path)

    assert total == 1


# ------------------------------------------------------------ 中断（Issue #61）


class FakeCancelEvent:
    """`is_set()` だけを持つ中断要求。

    run() は threading.Event を要求せず、`is_set()` を持つものなら何でも受ける
    （main.py は threading を import しない）。この偽物でテストが通ることが、その
    ダックタイピングを固定する。
    """

    def __init__(self):
        self._set = False

    def set(self):
        self._set = True

    def is_set(self):
        return self._set


def make_filled_workbook(tmp_path, rows):
    """(タイトル, 連載状況) の並びから、判定済みの行を含むブックを作る。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["タイトル", "連載状況"])
    for title, status in rows:
        ws.append([title, status])
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    return path


def test_run_stops_at_row_boundary_when_cancelled_and_saves_finished_rows(tmp_path):
    excel_path = make_workbook(tmp_path, ["A", "B", "C", "D", "E"])
    cancel_event = FakeCancelEvent()

    def judge(title):
        # 2行目を判定し終えた時点で中断を要求する（＝実行中に押された状況）。
        if title == "B":
            cancel_event.set()
        return WikipediaJudgment(query=title, status="完結", matched_title=title)

    calls = []

    def on_progress(index, total, title, status, matched_title, error=None):
        calls.append((index, total, title))

    with patch("comic_checker.main.wikipedia_client.judge", side_effect=judge):
        total = run(excel_path, on_progress=on_progress, cancel_event=cancel_event)

    # 戻り値は「処理対象の行数（5）」ではなく「実際に処理した行数（2）」。
    assert total == 2
    # 中断要求より後の行は判定にすら入らない（3行目のtotalは通知されない）。
    assert calls == [(1, 5, "A"), (2, 5, "B")]

    # save_every=5 に達する前に抜けているので、ループ後の保存が効いていないと
    # 2行とも消える。「破棄しない」の確認はここ。
    ws = _reload(excel_path).active
    assert read_row(ws, 2)["連載状況"] == "完結"
    assert read_row(ws, 3)["連載状況"] == "完結"
    # 未処理の行は空欄のまま＝次回 skip_filled=True の対象として残る。
    assert read_row(ws, 4)["連載状況"] is None
    assert read_row(ws, 5)["連載状況"] is None
    assert read_row(ws, 6)["連載状況"] is None


def test_run_without_cancel_event_processes_every_row(tmp_path):
    excel_path = make_workbook(tmp_path, ["A", "B", "C"])

    with patch(
        "comic_checker.main.wikipedia_client.judge",
        side_effect=lambda title: WikipediaJudgment(
            query=title, status="連載中", matched_title=title
        ),
    ):
        total = run(excel_path, cancel_event=None)

    assert total == 3
    ws = _reload(excel_path).active
    assert [read_row(ws, row)["連載状況"] for row in (2, 3, 4)] == ["連載中"] * 3


def test_cancelled_run_resumes_from_the_unprocessed_rows(tmp_path):
    excel_path = make_workbook(tmp_path, ["A", "B", "C", "D"])
    cancel_event = FakeCancelEvent()

    def judge_then_cancel(title):
        if title == "B":
            cancel_event.set()
        return WikipediaJudgment(query=title, status="完結", matched_title=title)

    with patch("comic_checker.main.wikipedia_client.judge", side_effect=judge_then_cancel):
        assert run(excel_path, cancel_event=cancel_event) == 2

    # 中断後に skip_filled=True で再実行すると、残り2行だけを続きから処理する。
    with patch(
        "comic_checker.main.wikipedia_client.judge",
        side_effect=lambda title: WikipediaJudgment(
            query=title, status="連載中", matched_title=title
        ),
    ) as resumed_judge:
        assert run(excel_path) == 2

    assert [call.args[0] for call in resumed_judge.call_args_list] == ["C", "D"]
    ws = _reload(excel_path).active
    # 1回目に処理した行は再判定されず、2回目の行だけが新しい値になる。
    assert [read_row(ws, row)["連載状況"] for row in (2, 3, 4, 5)] == [
        "完結",
        "完結",
        "連載中",
        "連載中",
    ]


def test_cancelled_rejudge_leaves_stale_rows_that_skip_filled_will_not_revisit(tmp_path):
    """skip_filled=False の実行を中断した場合の注意点（設計時に洗い出した挙動）。

    全件再判定を途中で止めると、まだ再判定していない行にも古い判定結果が残っている。
    それらは「判定済み」なので、次に skip_filled=True で実行しても**拾われない**。
    続きから再判定するには、もう一度「すべて再判定」を選ぶ必要がある（GUIは #32 の
    確認ダイアログでそれを選べる）。ここではその挙動を仕様として固定する。
    """
    excel_path = make_filled_workbook(
        tmp_path, [("A", "連載中"), ("B", "連載中"), ("C", "連載中")]
    )
    cancel_event = FakeCancelEvent()

    def judge_then_cancel(title):
        if title == "A":
            cancel_event.set()
        return WikipediaJudgment(query=title, status="完結", matched_title=title)

    with patch("comic_checker.main.wikipedia_client.judge", side_effect=judge_then_cancel):
        assert run(excel_path, skip_filled=False, cancel_event=cancel_event) == 1

    ws = _reload(excel_path).active
    # 再判定できたのは1行だけ。残り2行には古い「連載中」が残っている。
    assert [read_row(ws, row)["連載状況"] for row in (2, 3, 4)] == ["完結", "連載中", "連載中"]

    # skip_filled=True の再実行では、その古い2行は判定済み扱いで処理されない。
    with patch("comic_checker.main.wikipedia_client.judge") as never_judged:
        assert run(excel_path) == 0
    never_judged.assert_not_called()

    # 「すべて再判定」（skip_filled=False）を選べば、残りも新しい値に更新できる。
    with patch(
        "comic_checker.main.wikipedia_client.judge",
        side_effect=lambda title: WikipediaJudgment(
            query=title, status="完結", matched_title=title
        ),
    ):
        assert run(excel_path, skip_filled=False) == 3
    ws = _reload(excel_path).active
    assert [read_row(ws, row)["連載状況"] for row in (2, 3, 4)] == ["完結"] * 3


def _reload(excel_path):
    from openpyxl import load_workbook

    return load_workbook(excel_path)
