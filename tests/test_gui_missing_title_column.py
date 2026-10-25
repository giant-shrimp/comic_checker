"""保存したタイトル列名が Excel の1行目に無いときの案内のテスト（Issue #553）。名前はすべて架空。"""
import tkinter as tk
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from comic_checker import gui, user_config
from comic_checker.gui import (
    App,
    MISSING_SHEET_MESSAGE,
    MISSING_TITLE_COLUMN_MESSAGE,
    find_missing_title_column,
)
from comic_checker.main import RowCounts

EXPECTED_MESSAGE = (
    "列「存在しない列」が1行目に見つかりません。"
    "「設定」の画面のタイトル列名を確かめてください。"
)

ENTRY_POINTS = pytest.mark.parametrize("method_name", ["_on_run", "_on_next_release"])


@pytest.fixture
def app():
    root = tk.Tk()
    widget = App(master=root)
    widget.pack(fill="both", expand=True)
    root.update()
    yield widget
    widget._stop_elapsed()
    root.destroy()


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "list.xlsx"
    wb = Workbook()
    wb.active.title = "シートA"
    wb.active.append(["タイトル", "作者"])
    sheet_b = wb.create_sheet("シートB")
    sheet_b.append(["書名", "著者"])
    sheet_b.append(["存在しない列", "x"])
    wb.save(path)
    return str(path)


def save_values(**values):
    user_config.save_values(values)


def invoke(app, method_name):
    """指定の入口を呼び、(Thread のモック, count_rows のモック) を返す。"""
    with patch(
        "comic_checker.gui.count_rows", return_value=RowCounts(pending=3, filled=0)
    ) as count:
        with patch("comic_checker.gui.threading.Thread") as thread_cls:
            getattr(app, method_name)()
    return thread_cls, count


def log_messages(app):
    return [(level, message) for _time, level, message in app._log_entries]


def test_message_text_is_fixed():
    assert (
        MISSING_TITLE_COLUMN_MESSAGE.format(column="存在しない列") == EXPECTED_MESSAGE
    )


@ENTRY_POINTS
def test_missing_column_stops_before_starting(app, book, method_name):
    save_values(**{gui.KEY_TITLE_COLUMN: "存在しない列", gui.KEY_SHEET: "シートB"})
    app.excel_path.set(book)
    # 2行目に同じ名前があっても、1行目に無ければ止まる。
    thread_cls, count = invoke(app, method_name)
    thread_cls.assert_not_called()
    count.assert_not_called()
    assert ("ERROR", EXPECTED_MESSAGE) in log_messages(app)
    assert app.status_text.get() == EXPECTED_MESSAGE
    assert str(app.run_button.cget("state")) == "normal"
    assert str(app.next_button.cget("state")) == "normal"


@ENTRY_POINTS
def test_default_column_missing_in_saved_sheet_stops(app, book, method_name):
    save_values(**{gui.KEY_SHEET: "シートB"})
    app.excel_path.set(book)
    thread_cls, count = invoke(app, method_name)
    thread_cls.assert_not_called()
    count.assert_not_called()
    expected = MISSING_TITLE_COLUMN_MESSAGE.format(column="タイトル")
    assert ("ERROR", expected) in log_messages(app)
    assert app.status_text.get() == expected


@ENTRY_POINTS
def test_existing_column_starts_worker_with_title_column(app, book, method_name):
    save_values(**{gui.KEY_TITLE_COLUMN: "書名", gui.KEY_SHEET: "シートB"})
    app.excel_path.set(book)
    thread_cls, _count = invoke(app, method_name)
    thread_cls.assert_called_once()
    assert thread_cls.call_args.kwargs["kwargs"]["title_column"] == "書名"
    assert all(
        "1行目に見つかりません" not in m for _level, m in log_messages(app)
    )


@ENTRY_POINTS
def test_missing_sheet_shows_only_sheet_message(app, book, method_name):
    save_values(**{gui.KEY_TITLE_COLUMN: "存在しない列", gui.KEY_SHEET: "存在しないシート"})
    app.excel_path.set(book)
    thread_cls, _count = invoke(app, method_name)
    thread_cls.assert_not_called()
    messages = [m for _level, m in log_messages(app)]
    assert MISSING_SHEET_MESSAGE.format(sheet="存在しないシート") in messages
    assert all("1行目に見つかりません" not in m for m in messages)


@ENTRY_POINTS
def test_unopenable_file_falls_through_without_message(app, tmp_path, method_name):
    save_values(**{gui.KEY_TITLE_COLUMN: "存在しない列"})
    app.excel_path.set(str(tmp_path / "no_such_file.xlsx"))
    thread_cls, _count = invoke(app, method_name)
    thread_cls.assert_called_once()
    assert all(
        "1行目に見つかりません" not in m for _level, m in log_messages(app)
    )


def make_book(rows, sheetnames=("シートA",)):
    wb = MagicMock()
    wb.sheetnames = list(sheetnames)
    ws = MagicMock()
    ws.iter_rows.return_value = iter(rows)
    wb.active = ws
    wb.__getitem__.return_value = ws
    return wb, ws


def test_find_missing_title_column_looks_only_at_first_row():
    wb, ws = make_book([("書名", "著者")])
    with patch("comic_checker.gui.load_workbook", return_value=wb):
        assert find_missing_title_column("/tmp/l.xlsx", "存在しない列", "") is True
    ws.iter_rows.assert_called_once_with(min_row=1, max_row=1, values_only=True)


def test_find_missing_title_column_ignores_second_row(tmp_path):
    path = tmp_path / "second.xlsx"
    wb = Workbook()
    wb.active.append(["書名"])
    wb.active.append(["存在しない列"])
    wb.save(path)
    assert find_missing_title_column(str(path), "存在しない列", None) is True
    assert find_missing_title_column(str(path), "書名", None) is False


def test_find_missing_title_column_requires_exact_match(tmp_path):
    path = tmp_path / "exact.xlsx"
    wb = Workbook()
    wb.active.append(["書名 "])
    wb.save(path)
    assert find_missing_title_column(str(path), "書名", None) is True
    assert find_missing_title_column(str(path), "書名 ", None) is False


def test_find_missing_title_column_empty_sheet_is_true():
    wb, _ws = make_book([])
    with patch("comic_checker.gui.load_workbook", return_value=wb):
        assert find_missing_title_column("/tmp/l.xlsx", "書名", "シートA") is True


def test_find_missing_title_column_opens_read_only_and_closes_once():
    wb, _ws = make_book([("書名",)])
    with patch("comic_checker.gui.load_workbook", return_value=wb) as loader:
        assert find_missing_title_column("/tmp/l.xlsx", "書名", "シートA") is False
    loader.assert_called_once_with("/tmp/l.xlsx", read_only=True)
    wb.close.assert_called_once_with()


def test_find_missing_title_column_missing_sheet_is_false_and_closes():
    wb, _ws = make_book([("書名",)])
    with patch("comic_checker.gui.load_workbook", return_value=wb):
        assert (
            find_missing_title_column("/tmp/l.xlsx", "存在しない列", "無いシート")
            is False
        )
    wb.close.assert_called_once_with()


def test_find_missing_title_column_returns_false_when_open_fails():
    with patch("comic_checker.gui.load_workbook", side_effect=OSError("x")):
        assert find_missing_title_column("/tmp/l.xlsx", "書名", "シートA") is False
