"""Wikipedia が記事なしの行だけ BOOK☆WALKER で判定し直すことのテスト。作品名はすべて架空。"""
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from comic_checker.bookwalker_client import BookwalkerResult
from comic_checker.main import BOOKWALKER_SOURCE, WIKIPEDIA_SOURCE, parse_args, run
from comic_checker.wikipedia_client import WikipediaJudgment


def make_workbook(tmp_path, titles):
    wb = Workbook()
    ws = wb.active
    ws.append(["タイトル", "連載状況"])
    for title in titles:
        ws.append([title, None])
    path = tmp_path / "test.xlsx"
    wb.save(path)
    return path


def read_row(ws, row):
    headers = [cell.value for cell in ws[1]]
    return {header: ws.cell(row=row, column=col).value for col, header in enumerate(headers, start=1)}


def run_with(tmp_path, wiki_result, bw_result=None, **kwargs):
    excel_path = make_workbook(tmp_path, ["作品A"])
    with patch("comic_checker.main.wikipedia_client.judge", return_value=wiki_result):
        with patch("comic_checker.main.bookwalker_client.lookup", return_value=bw_result) as lookup:
            run(excel_path, **kwargs)
    return read_row(load_workbook(excel_path).active, 2), lookup


def not_found():
    return WikipediaJudgment(query="作品A", status="記事なし")


def test_not_found_row_takes_the_bookwalker_completed_answer(tmp_path):
    bw = BookwalkerResult(
        "完結", "作品A（架空レーベル）", "https://bookwalker.jp/series/1/list/", "完結の札あり（シリーズ8冊）"
    )
    row, lookup = run_with(tmp_path, not_found(), bw)
    lookup.assert_called_once_with("作品A")
    assert row["連載状況"] == "完結"
    assert row["判定方法"] == "BOOK☆WALKER"
    assert row["判定根拠"] == "完結の札あり（シリーズ8冊）"
    assert row["参照タイトル"] == "作品A（架空レーベル）"
    assert row["参照URL"] == "https://bookwalker.jp/series/1/list/"


def test_not_found_row_takes_the_bookwalker_ongoing_answer(tmp_path):
    bw = BookwalkerResult("連載中", "作品A（架空レーベル）", "https://bookwalker.jp/series/2/list/", "完結の札なし")
    row, _ = run_with(tmp_path, not_found(), bw)
    assert row["連載状況"] == "連載中"
    assert row["判定方法"] == BOOKWALKER_SOURCE


def test_undecided_bookwalker_answers_keep_not_found(tmp_path):
    for status in ("一致なし", "複数一致", "エラー"):
        row, _ = run_with(tmp_path, not_found(), BookwalkerResult(status, evidence="取得できませんでした（HTTP 403）"))
        assert row["連載状況"] == "記事なし"
        assert row["判定方法"] == WIKIPEDIA_SOURCE
        assert row["判定根拠"] in (None, "")


def test_wikipedia_answers_do_not_ask_bookwalker(tmp_path):
    for status in ("完結", "連載中", "判定不能"):
        wiki = WikipediaJudgment(query="作品A", status=status, matched_title="作品A")
        row, lookup = run_with(tmp_path, wiki, BookwalkerResult("完結"))
        lookup.assert_not_called()
        assert row["連載状況"] == status
        assert row["判定方法"] == WIKIPEDIA_SOURCE


def test_bookwalker_can_be_turned_off(tmp_path):
    row, lookup = run_with(tmp_path, not_found(), BookwalkerResult("完結"), use_bookwalker=False)
    lookup.assert_not_called()
    assert row["連載状況"] == "記事なし"


def test_not_found_row_stays_not_found_when_bookwalker_cannot_be_reached(tmp_path):
    # lookup はモックしない。conftest が取得関数を「取得できない」に差し替えている。
    excel_path = make_workbook(tmp_path, ["作品A"])
    with patch("comic_checker.main.wikipedia_client.judge", return_value=not_found()):
        run(excel_path)
    row = read_row(load_workbook(excel_path).active, 2)
    assert row["連載状況"] == "記事なし"
    assert row["判定方法"] == WIKIPEDIA_SOURCE


def test_cli_uses_bookwalker_by_default_and_can_turn_it_off():
    assert parse_args(["list.xlsx"]).use_bookwalker is True
    assert parse_args(["list.xlsx", "--no-bookwalker"]).use_bookwalker is False
